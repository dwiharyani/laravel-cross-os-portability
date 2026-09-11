#!/usr/bin/env python3

"""
06_filter_eligible_laravel.py

FINAL Laravel eligibility screening.

Data sources
------------
1. github_enriched.csv
   - GitHub repository metadata
   - pushed_at
   - default_branch
   - latest_sha
   - archived
   - fork
   - license
   - language

2. verified_laravel_apps.csv
   - Laravel verification
   - Laravel classification
   - artisan
   - app_dir
   - bootstrap_dir
   - config_dir
   - routes_dir
   - tests_dir

The two datasets are merged using repo_full_name.

Deep repository checks are performed only for
Confirmed Laravel Application repositories.

Checks
------
Activity:
    pushed_at

Testing:
    tests/
    phpunit.xml
    phpunit.xml.dist
    Composer test scripts

Dependencies:
    composer.json
    composer.lock

Laravel:
    laravel/framework constraint

PHP:
    PHP constraint

Docker:
    Dockerfile
    docker-compose.yml
    docker-compose.yaml
    compose.yml
    compose.yaml

Docker is recorded and, by default, excluded because the
planned experiment evaluates native cross-OS portability.
"""

from __future__ import annotations

import argparse
import base64
import json
import os
from datetime import datetime, timezone

import pandas as pd
import requests


# ============================================================
# PATHS
# ============================================================

ROOT = os.path.abspath(
    os.path.join(
        os.path.dirname(__file__),
        "..",
    )
)

GITHUB_FILE = os.path.join(
    ROOT,
    "data/interim/github_enriched.csv",
)

LARAVEL_FILE = os.path.join(
    ROOT,
    "data/interim/verified_laravel_apps.csv",
)

OUTPUT_CSV = os.path.join(
    ROOT,
    "data/interim/eligible_laravel_apps.csv",
)

OUTPUT_XLSX = os.path.join(
    ROOT,
    "data/interim/laravel_eligibility_results.xlsx",
)


# ============================================================
# CONFIGURATION
# ============================================================

RECENT_MONTHS = 24

REQUIRE_TESTS_DIRECTORY = True

REQUIRE_PHPUNIT_OR_COMPOSER_TEST = True

EXCLUDE_DOCKER = True


DOCKER_FILES = [
    "Dockerfile",
    "docker-compose.yml",
    "docker-compose.yaml",
    "compose.yml",
    "compose.yaml",
]


PHPUNIT_FILES = [
    "phpunit.xml",
    "phpunit.xml.dist",
]


# ============================================================
# GITHUB AUTHENTICATION
# ============================================================

def load_github_token():

    token = os.getenv(
        "GITHUB_TOKEN"
    )

    if token:
        return token

    env_file = os.path.join(
        ROOT,
        ".env",
    )

    if os.path.exists(env_file):

        with open(
            env_file,
            "r",
            encoding="utf-8",
        ) as f:

            for line in f:

                line = line.strip()

                if (
                    not line
                    or line.startswith("#")
                ):
                    continue

                if line.startswith(
                    "GITHUB_TOKEN="
                ):

                    return (
                        line.split(
                            "=",
                            1,
                        )[1]
                        .strip()
                        .strip('"')
                        .strip("'")
                    )

    return None


TOKEN = load_github_token()


HEADERS = {
    "Accept":
        "application/vnd.github+json",

    "X-GitHub-Api-Version":
        "2022-11-28",
}


if TOKEN:

    HEADERS[
        "Authorization"
    ] = f"Bearer {TOKEN}"


# ============================================================
# GITHUB API
# ============================================================

SESSION = requests.Session()

SESSION.headers.update(
    HEADERS
)


def github_get(
    url,
    timeout=30,
):

    try:

        response = SESSION.get(
            url,
            timeout=timeout,
        )

        if response.status_code == 200:

            return response.json()

        return {
            "_error": True,
            "_status": response.status_code,
        }

    except Exception as e:

        return {
            "_error": True,
            "_exception": repr(e),
        }


def repo_contents(
    repo,
    path,
    ref,
):

    url = (
        "https://api.github.com/repos/"
        f"{repo}/contents/{path}"
    )

    response = github_get(
        f"{url}?ref={ref}"
    )

    return response


def github_file_exists(
    repo,
    filename,
    ref,
):

    data = repo_contents(
        repo,
        filename,
        ref,
    )

    if isinstance(
        data,
        dict,
    ) and data.get("_error"):

        return False

    return isinstance(
        data,
        dict,
    )


def github_directory_exists(
    repo,
    dirname,
    ref,
):

    data = repo_contents(
        repo,
        dirname,
        ref,
    )

    if isinstance(
        data,
        dict,
    ) and data.get("_error"):

        return False

    return isinstance(
        data,
        list,
    )


def github_file_text(
    repo,
    filename,
    ref,
):

    data = repo_contents(
        repo,
        filename,
        ref,
    )

    if not isinstance(
        data,
        dict,
    ):

        return ""

    if data.get("_error"):

        return ""

    encoded = data.get(
        "content"
    )

    if not encoded:

        return ""

    try:

        return base64.b64decode(
            encoded
        ).decode(
            "utf-8",
            errors="ignore",
        )

    except Exception:

        return ""


# ============================================================
# HELPERS
# ============================================================

def clean(value):

    if value is None:

        return ""

    if pd.isna(value):

        return ""

    return str(
        value
    ).strip()


def as_bool(value):

    if isinstance(
        value,
        bool,
    ):

        return value

    return clean(
        value
    ).lower() in {
        "true",
        "1",
        "yes",
    }


def parse_datetime(value):

    value = clean(
        value
    )

    if not value:

        return None

    try:

        return pd.to_datetime(
            value,
            utc=True,
        )

    except Exception:

        return None


# ============================================================
# LOAD DATA
# ============================================================

def load_input_data():

    if not os.path.exists(
        GITHUB_FILE
    ):

        raise FileNotFoundError(
            GITHUB_FILE
        )

    if not os.path.exists(
        LARAVEL_FILE
    ):

        raise FileNotFoundError(
            LARAVEL_FILE
        )

    github = pd.read_csv(
        GITHUB_FILE
    )

    laravel = pd.read_csv(
        LARAVEL_FILE
    )

    return github, laravel


# ============================================================
# MERGE DATA
# ============================================================

def prepare_confirmed_laravel(
    github,
    laravel,
):

    confirmed = laravel[
        laravel[
            "laravel_classification"
        ]
        ==
        "Confirmed Laravel Application"
    ].copy()

    print(
        "Confirmed Laravel applications:",
        len(confirmed),
    )

    merged = confirmed.merge(
        github,
        on="repo_full_name",
        how="left",
        suffixes=(
            "_laravel",
            "_github",
        ),
    )

    print(
        "Merged rows:",
        len(merged),
    )

    return merged


# ============================================================
# ACTIVITY
# ============================================================

def activity_check(
    pushed_at,
):

    pushed = parse_datetime(
        pushed_at
    )

    if pushed is None:

        return (
            False,
            "",
        )

    cutoff = (
        pd.Timestamp.now(
            tz="UTC"
        )
        -
        pd.DateOffset(
            months=RECENT_MONTHS
        )
    )

    return (
        pushed >= cutoff,
        pushed.isoformat(),
    )


# ============================================================
# COMPOSER
# ============================================================

def composer_analysis(
    repo,
    ref,
):

    result = {
        "composer_json_present": False,
        "composer_lock_present": False,
        "composer_test_script_present": False,
        "composer_test_scripts": "",
        "php_version_constraint": "",
        "laravel_version_constraint": "",
    }

    composer_text = github_file_text(
        repo,
        "composer.json",
        ref,
    )

    if not composer_text:

        return result

    result[
        "composer_json_present"
    ] = True

    result[
        "composer_lock_present"
    ] = github_file_exists(
        repo,
        "composer.lock",
        ref,
    )

    try:

        data = json.loads(
            composer_text
        )

    except Exception:

        return result

    require = data.get(
        "require",
        {},
    )

    if isinstance(
        require,
        dict,
    ):

        result[
            "php_version_constraint"
        ] = clean(
            require.get(
                "php",
                "",
            )
        )

        result[
            "laravel_version_constraint"
        ] = clean(
            require.get(
                "laravel/framework",
                "",
            )
        )

    scripts = data.get(
        "scripts",
        {},
    )

    test_scripts = []

    if isinstance(
        scripts,
        dict,
    ):

        for name, command in scripts.items():

            name_lower = clean(
                name
            ).lower()

            if (
                "test" in name_lower
                or name_lower in {
                    "phpunit",
                    "pest",
                }
            ):

                test_scripts.append(
                    f"{name}: {command}"
                )

    result[
        "composer_test_scripts"
    ] = " | ".join(
        test_scripts
    )

    result[
        "composer_test_script_present"
    ] = bool(
        test_scripts
    )

    return result


# ============================================================
# TESTING
# ============================================================

def testing_analysis(
    repo,
    ref,
    existing_tests_dir,
):

    result = {
        "tests_directory_present": bool(
            existing_tests_dir
        ),
        "phpunit_config_present": False,
        "phpunit_config_file": "",
    }

    found = []

    for filename in PHPUNIT_FILES:

        if github_file_exists(
            repo,
            filename,
            ref,
        ):

            found.append(
                filename
            )

    result[
        "phpunit_config_present"
    ] = bool(
        found
    )

    result[
        "phpunit_config_file"
    ] = "; ".join(
        found
    )

    return result


# ============================================================
# DOCKER
# ============================================================

def docker_analysis(
    repo,
    ref,
):

    found = []

    for filename in DOCKER_FILES:

        if github_file_exists(
            repo,
            filename,
            ref,
        ):

            found.append(
                filename
            )

    return {
        "dockerfile_present":
            "Dockerfile" in found,

        "compose_present":
            any(
                x != "Dockerfile"
                for x in found
            ),

        "compose_files":
            "; ".join(
                [
                    x
                    for x in found
                    if x != "Dockerfile"
                ]
            ),

        "docker_present":
            bool(found),
    }


# ============================================================
# SCREEN ONE REPOSITORY
# ============================================================

def screen_repository(
    row,
    index,
    total,
):

    repo = clean(
        row.get(
            "repo_full_name",
            "",
        )
    )

    branch = clean(
        row.get(
            "default_branch",
            "",
        )
    )

    sha = clean(
        row.get(
            "latest_sha",
            "",
        )
    )

    ref = sha or branch or "main"

    print(
        f"[{index}/{total}] {repo}"
    )

    result = {}

    # --------------------------------------------------------
    # Basic metadata
    # --------------------------------------------------------

    result[
        "project_id"
    ] = clean(
        row.get(
            "project_id_laravel",
            "",
        )
    )

    result[
        "repo_full_name"
    ] = repo

    result[
        "url"
    ] = clean(
        row.get(
            "url_github",
            row.get(
                "url_laravel",
                "",
            ),
        )
    )

    result[
        "default_branch"
    ] = branch

    result[
        "latest_sha"
    ] = sha

    result[
        "stars"
    ] = clean(
        row.get(
            "stars_github",
            "",
        )
    )

    result[
        "license"
    ] = clean(
        row.get(
            "license",
            "",
        )
    )

    result[
        "created_at"
    ] = clean(
        row.get(
            "created_at",
            "",
        )
    )

    result[
        "updated_at"
    ] = clean(
        row.get(
            "updated_at",
            "",
        )
    )

    result[
        "pushed_at"
    ] = clean(
        row.get(
            "pushed_at",
            "",
        )
    )

    result[
        "size_kb"
    ] = clean(
        row.get(
            "size_kb",
            "",
        )
    )

    result[
        "archived"
    ] = as_bool(
        row.get(
            "archived",
            False,
        )
    )

    result[
        "fork"
    ] = as_bool(
        row.get(
            "fork",
            False,
        )
    )

    result[
        "disabled"
    ] = as_bool(
        row.get(
            "disabled",
            False,
        )
    )

    # --------------------------------------------------------
    # Existing Laravel verification
    # --------------------------------------------------------

    result[
        "laravel_classification"
    ] = clean(
        row.get(
            "laravel_classification",
            "",
        )
    )

    result[
        "artisan"
    ] = as_bool(
        row.get(
            "artisan",
            False,
        )
    )

    result[
        "app_dir"
    ] = as_bool(
        row.get(
            "app_dir",
            False,
        )
    )

    result[
        "bootstrap_dir"
    ] = as_bool(
        row.get(
            "bootstrap_dir",
            False,
        )
    )

    result[
        "config_dir"
    ] = as_bool(
        row.get(
            "config_dir",
            False,
        )
    )

    result[
        "routes_dir"
    ] = as_bool(
        row.get(
            "routes_dir",
            False,
        )
    )

    result[
        "tests_dir_initial"
    ] = as_bool(
        row.get(
            "tests_dir",
            False,
        )
    )

    # --------------------------------------------------------
    # Activity
    # --------------------------------------------------------

    (
        recent,
        parsed_push,
    ) = activity_check(
        row.get(
            "pushed_at",
            "",
        )
    )

    result[
        "recent_activity"
    ] = recent

    result[
        "parsed_pushed_at"
    ] = parsed_push

    # --------------------------------------------------------
    # Testing
    # --------------------------------------------------------

    testing = testing_analysis(
        repo,
        ref,
        result[
            "tests_dir_initial"
        ],
    )

    result.update(
        testing
    )

    # --------------------------------------------------------
    # Composer
    # --------------------------------------------------------

    composer = composer_analysis(
        repo,
        ref,
    )

    result.update(
        composer
    )

    # --------------------------------------------------------
    # Docker
    # --------------------------------------------------------

    docker = docker_analysis(
        repo,
        ref,
    )

    result.update(
        docker
    )

    # --------------------------------------------------------
    # Eligibility
    # --------------------------------------------------------

    reasons = []

    if result[
        "laravel_classification"
    ] != "Confirmed Laravel Application":

        reasons.append(
            "Not confirmed Laravel application"
        )

    if not result[
        "recent_activity"
    ]:

        reasons.append(
            "Not recently active"
        )

    if (
        REQUIRE_TESTS_DIRECTORY
        and
        not result[
            "tests_directory_present"
        ]
    ):

        reasons.append(
            "No tests directory"
        )

    if REQUIRE_PHPUNIT_OR_COMPOSER_TEST:

        has_testing_config = (
            result[
                "phpunit_config_present"
            ]
            or
            result[
                "composer_test_script_present"
            ]
        )

        if not has_testing_config:

            reasons.append(
                "No PHPUnit configuration "
                "or Composer test script"
            )

    if not result[
        "composer_json_present"
    ]:

        reasons.append(
            "No composer.json"
        )

    if (
        EXCLUDE_DOCKER
        and
        result[
            "docker_present"
        ]
    ):

        reasons.append(
            "Docker/Compose present"
        )

    result[
        "eligible"
    ] = not bool(
        reasons
    )

    result[
        "exclusion_reason"
    ] = (
        "Eligible"
        if result["eligible"]
        else "; ".join(
            reasons
        )
    )

    result[
        "screening_timestamp"
    ] = datetime.now(
        timezone.utc
    ).isoformat()

    return result


# ============================================================
# EXCEL
# ============================================================

def save_excel(
    results,
    output_path,
    configuration,
):

    from openpyxl import load_workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
    ) as writer:

        results.to_excel(
            writer,
            sheet_name="All Results",
            index=False,
        )

        results[
            results["eligible"]
        ].to_excel(
            writer,
            sheet_name="Eligible",
            index=False,
        )

        results[
            ~results["eligible"]
        ].to_excel(
            writer,
            sheet_name="Excluded",
            index=False,
        )

        (
            results[
                "exclusion_reason"
            ]
            .value_counts()
            .rename_axis(
                "exclusion_reason"
            )
            .reset_index(
                name="count"
            )
            .to_excel(
                writer,
                sheet_name="Exclusion Reasons",
                index=False,
            )
        )

        summary = pd.DataFrame(
            [
                {
                    "metric":
                        "Total screened",
                    "value":
                        len(results),
                },
                {
                    "metric":
                        "Eligible",
                    "value":
                        int(
                            results[
                                "eligible"
                            ].sum()
                        ),
                },
                {
                    "metric":
                        "Excluded",
                    "value":
                        int(
                            (
                                ~results[
                                    "eligible"
                                ]
                            ).sum()
                        ),
                },
                {
                    "metric":
                        "Docker detected",
                    "value":
                        int(
                            results[
                                "docker_present"
                            ].sum()
                        ),
                },
                {
                    "metric":
                        "Recently active",
                    "value":
                        int(
                            results[
                                "recent_activity"
                            ].sum()
                        ),
                },
                {
                    "metric":
                        "Tests directory",
                    "value":
                        int(
                            results[
                                "tests_directory_present"
                            ].sum()
                        ),
                },
                {
                    "metric":
                        "PHPUnit config",
                    "value":
                        int(
                            results[
                                "phpunit_config_present"
                            ].sum()
                        ),
                },
                {
                    "metric":
                        "Composer test script",
                    "value":
                        int(
                            results[
                                "composer_test_script_present"
                            ].sum()
                        ),
                },
            ]
        )

        summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
        )

        pd.DataFrame(
            [
                {
                    "parameter": key,
                    "value": value,
                }
                for key, value
                in configuration.items()
            ]
        ).to_excel(
            writer,
            sheet_name="Configuration",
            index=False,
        )

    # Formatting
    wb = load_workbook(
        output_path
    )

    for ws in wb.worksheets:

        ws.freeze_panes = "A2"

        ws.auto_filter.ref = (
            ws.dimensions
        )

        for cell in ws[1]:

            cell.font = Font(
                bold=True
            )

        for column_cells in ws.columns:

            max_len = 0

            col = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_len = max(
                        max_len,
                        len(
                            str(
                                cell.value
                            )
                        ),
                    )

            ws.column_dimensions[
                col
            ].width = min(
                max(
                    max_len + 2,
                    12,
                ),
                60,
            )

    wb.save(
        output_path
    )


# ============================================================
# MAIN
# ============================================================

def main():

    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--limit",
        type=int,
        default=None,
    )

    args = parser.parse_args()

    print("=" * 70)
    print(
        "FINAL LARAVEL ELIGIBILITY SCREENING"
    )
    print("=" * 70)

    print(
        "GitHub metadata:",
        GITHUB_FILE,
    )

    print(
        "Laravel verification:",
        LARAVEL_FILE,
    )

    print()

    github, laravel = (
        load_input_data()
    )

    merged = prepare_confirmed_laravel(
        github,
        laravel,
    )

    if args.limit:

        merged = merged.head(
            args.limit
        ).copy()

        print(
            "PILOT LIMIT:",
            len(merged),
        )

    results = []

    total = len(
        merged
    )

    for i, (_, row) in enumerate(
        merged.iterrows(),
        start=1,
    ):

        try:

            result = screen_repository(
                row,
                i,
                total,
            )

            results.append(
                result
            )

        except Exception as e:

            repo = clean(
                row.get(
                    "repo_full_name",
                    "",
                )
            )

            print(
                "ERROR:",
                repo,
                repr(e),
            )

            results.append(
                {
                    "repo_full_name":
                        repo,

                    "eligible":
                        False,

                    "exclusion_reason":
                        "Screening error: "
                        + repr(e),
                }
            )

    results = pd.DataFrame(
        results
    )

    # --------------------------------------------------------
    # Save CSV
    # --------------------------------------------------------

    os.makedirs(
        os.path.dirname(
            OUTPUT_CSV
        ),
        exist_ok=True,
    )

    results.to_csv(
        OUTPUT_CSV,
        index=False,
    )

    # --------------------------------------------------------
    # Save Excel
    # --------------------------------------------------------

    configuration = {
        "recent_activity_months":
            RECENT_MONTHS,

        "require_tests_directory":
            REQUIRE_TESTS_DIRECTORY,

        "require_phpunit_or_composer_test":
            REQUIRE_PHPUNIT_OR_COMPOSER_TEST,

        "exclude_docker":
            EXCLUDE_DOCKER,

        "docker_files":
            "; ".join(
                DOCKER_FILES
            ),

        "phpunit_files":
            "; ".join(
                PHPUNIT_FILES
            ),

        "screening_timestamp":
            datetime.now(
                timezone.utc
            ).isoformat(),
    }

    save_excel(
        results,
        OUTPUT_XLSX,
        configuration,
    )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print(
        "ELIGIBILITY SUMMARY"
    )
    print("=" * 70)

    print(
        "Total screened:",
        len(results),
    )

    print(
        "Eligible:",
        int(
            results[
                "eligible"
            ].sum()
        ),
    )

    print(
        "Not eligible:",
        int(
            (
                ~results[
                    "eligible"
                ]
            ).sum()
        ),
    )

    print()
    print(
        "EXCLUSION REASONS"
    )

    print(
        results[
            "exclusion_reason"
        ]
        .value_counts()
        .to_string()
    )

    print()
    print(
        "OUTPUTS"
    )

    print(
        "CSV:",
        OUTPUT_CSV,
    )

    print(
        "Excel:",
        OUTPUT_XLSX,
    )

    print("=" * 70)


if __name__ == "__main__":
    main()
